#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
exportar_json.py — Convierte comparativa_fondeadoras.xlsx en fondeadoras.json,
el fichero que lee la app del movil.

Uso:
    python exportar_json.py
    python exportar_json.py --excel ruta\\comparativa_fondeadoras.xlsx --salida docs\\fondeadoras.json

Requiere:  pip install openpyxl

El JSON tiene tres bloques:
    meta    -> titulo, fechas, avisos, textos del simulador y el metodo
    campos  -> ESQUEMA de cada columna (tipo, grupo, si es filtrable/ordenable)
    planes  -> una fila por plan

La app construye los filtros, el orden, el selector de metricas, el comparador,
la ficha y el CSV LEYENDO el bloque "campos". Si anades una columna al Excel y
la declaras en MAPA_COLUMNAS, aparece sola en la app: no hay que tocar el HTML.
Si no la declaras, el script la exporta igualmente como texto y te avisa.
"""

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl.  Instalalo con:  pip install openpyxl")


# --------------------------------------------------------------------------
# 1. ESQUEMA: cabecera del Excel -> definicion del campo
# --------------------------------------------------------------------------
# tipo:      numero | texto | enum | fecha
# grupo:     seccion de la ficha y del comparador (el orden lo fija ORDEN_GRUPOS)
# filtrable: los enum salen como facetas con casillas; los numero, como rangos
# ordenable: aparece en la hoja "Ordenar por"
MAPA_COLUMNAS = {
    "Puesto": dict(key="puesto", tipo="numero", grupo="Identidad", ordenable=True),
    "Firma": dict(key="firma", tipo="enum", grupo="Identidad", filtrable=True, ordenable=True),
    "Plan": dict(key="plan", tipo="texto", grupo="Identidad", ordenable=True),
    "Estado y motivo de exclusion": dict(key="estado", tipo="texto", grupo="Identidad"),

    "Precio tarifa": dict(key="precio_tarifa", tipo="numero", unidad="USD", grupo="Precio",
                          filtrable=True, ordenable=True),
    "Precio con promo": dict(key="precio_promo", tipo="numero", unidad="USD", grupo="Precio",
                             filtrable=True, ordenable=True),
    "Cupon": dict(key="cupon", tipo="texto", grupo="Precio"),
    "Promo caduca": dict(key="promo_caduca", tipo="texto", grupo="Precio"),
    "Tipo precio": dict(key="tipo_precio", tipo="enum", grupo="Precio",
                        filtrable=True, ordenable=True),
    "Activacion": dict(key="activacion", tipo="numero", unidad="USD", grupo="Precio",
                       filtrable=True, ordenable=True),

    "Objetivo": dict(key="objetivo", tipo="numero", unidad="USD", grupo="Reglas",
                     filtrable=True, ordenable=True),
    "Drawdown en evaluacion": dict(key="dd_evaluacion", tipo="enum", grupo="Reglas",
                                   filtrable=True, ordenable=True),
    "Drawdown en cuenta fondeada": dict(key="dd_fondeada", tipo="enum", grupo="Reglas",
                                        filtrable=True),
    "Cambia al fondearse": dict(key="cambia_fondearse", tipo="enum", grupo="Reglas",
                                filtrable=True),
    "Drawdown": dict(key="drawdown", tipo="numero", unidad="USD", grupo="Reglas",
                     filtrable=True, ordenable=True),
    "Perdida diaria": dict(key="perdida_diaria", tipo="numero", unidad="USD", grupo="Reglas",
                           filtrable=True, ordenable=True),
    "Consistencia %": dict(key="consistencia_pct", tipo="numero", unidad="%", grupo="Reglas",
                           filtrable=True, ordenable=True),
    "Consistencia: donde aplica": dict(key="consistencia_donde", tipo="enum", grupo="Reglas",
                                       filtrable=True),

    "Min. diario para cobrar": dict(key="min_diario_cobro", tipo="numero", unidad="USD",
                                    grupo="Cobros", filtrable=True, ordenable=True),
    "Techo total de cobros": dict(key="techo_cobros", tipo="texto", grupo="Cobros"),
    "Dias min.": dict(key="dias_min", tipo="numero", unidad="dias", grupo="Cobros",
                      filtrable=True, ordenable=True),
    "Reparto": dict(key="reparto", tipo="enum", grupo="Cobros", filtrable=True, ordenable=True),
    "Buffer cobro": dict(key="buffer_cobro", tipo="numero", unidad="USD", grupo="Cobros",
                         filtrable=True, ordenable=True),

    "Broker": dict(key="broker", tipo="texto", grupo="Plataforma"),
    "TradingView": dict(key="tradingview", tipo="enum", grupo="Plataforma", filtrable=True),
    "NinjaTrader": dict(key="ninjatrader", tipo="enum", grupo="Plataforma", filtrable=True),

    "Capital propio en riesgo": dict(key="capital_riesgo", tipo="enum", grupo="Fiabilidad",
                                     filtrable=True),
    "Verificado en fuente oficial": dict(key="verificado", tipo="enum", grupo="Fiabilidad",
                                         filtrable=True, ordenable=True),
    "Fecha consulta": dict(key="fecha_consulta", tipo="fecha", grupo="Fiabilidad", ordenable=True),
    "QUE LE LASTRA LA NOTA": dict(key="lastra_nota", tipo="texto", grupo="Fiabilidad"),
    "Fiabilidad: por que esa nota": dict(key="fiabilidad_por_que", tipo="texto", grupo="Fiabilidad"),
}

# Campo que no esta en el Excel: se extrae del "73/100 — ..." de fiabilidad_por_que
# para poder ORDENAR y FILTRAR por la nota, cosa que con el texto no se puede.
CAMPO_NOTA = dict(key="fiabilidad_nota", label="Nota de fiabilidad", tipo="numero",
                  unidad="/100", grupo="Fiabilidad", filtrable=True, ordenable=True)

ORDEN_GRUPOS = ["Identidad", "Precio", "Reglas", "Cobros", "Plataforma", "Fiabilidad"]

COSTE_NOTA = ("COSTE TARIFA es el precio oficial sin descuentos: es lo comparable en el tiempo "
              "y es con lo que puntua el ranking. COSTE CON PROMO es lo que pagarias HOY con la "
              "oferta vigente. Cuando caduque una promocion, esa columna deja de ser cierta: "
              "mira siempre la fecha de caducidad.")
COSTE_AVISO = ("Ninguno de los dos costes incluye resets ni renovaciones tras un incumplimiento, "
               "que en la practica es la partida que mas se dispara.")


# --------------------------------------------------------------------------
# 2. Utilidades
# --------------------------------------------------------------------------
def slug(texto):
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    t = re.sub(r"[^a-zA-Z0-9]+", "-", t).strip("-").lower()
    return t or "x"


def limpiar(valor):
    """Normaliza una celda: cadena vacia o solo espacios -> None."""
    if valor is None:
        return None
    if isinstance(valor, str):
        v = valor.strip()
        return v if v else None
    if isinstance(valor, (dt.datetime, dt.date)):
        return valor.strftime("%Y-%m-%d")
    return valor


def a_numero(valor):
    """int si el numero es entero, float si no, None si no hay numero."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        n = round(float(valor), 2)
    else:
        m = re.search(r"-?\d+(?:[.,]\d+)?", str(valor))
        if not m:
            return None
        n = round(float(m.group(0).replace(",", ".")), 2)
    return int(n) if n == int(n) else n


def fila_vacia(fila):
    return all(limpiar(c) is None for c in fila)


# --------------------------------------------------------------------------
# 3. Hoja Comparativa
# --------------------------------------------------------------------------
def leer_comparativa(hoja, avisos):
    filas = list(hoja.iter_rows(values_only=True))

    idx = next((i for i, f in enumerate(filas) if limpiar(f[0]) == "Puesto"), None)
    if idx is None:
        sys.exit("No encuentro la fila de cabecera (la que empieza por 'Puesto') en 'Comparativa'.")

    cabecera = [limpiar(c) for c in filas[idx]]
    subtitulo = limpiar(filas[1][0]) if len(filas) > 1 else None

    # La leyenda de colores: ultima fila larga sin firma en la segunda columna
    leyenda = None
    for f in reversed(filas[idx + 1:]):
        t = limpiar(f[0])
        if t and limpiar(f[1]) is None and len(str(t)) > 40:
            leyenda = t
            break

    campos, desconocidas = [], []
    for col, nombre in enumerate(cabecera):
        if nombre is None:
            continue
        if nombre in MAPA_COLUMNAS:
            d = dict(MAPA_COLUMNAS[nombre])
        else:
            desconocidas.append(nombre)
            d = dict(key=slug(nombre).replace("-", "_"), tipo="texto", grupo="Otros",
                     filtrable=True)
        d.setdefault("label", nombre)
        d["_col"] = col
        campos.append(d)

    if desconocidas:
        avisos.append("Columnas nuevas, exportadas como texto en el grupo 'Otros': "
                      + ", ".join(desconocidas))

    faltan = [n for n in MAPA_COLUMNAS if n not in cabecera]
    if faltan:
        avisos.append("Columnas declaradas que ya NO estan en el Excel: " + ", ".join(faltan))

    planes = []
    for fila in filas[idx + 1:]:
        if fila_vacia(fila) or limpiar(fila[1]) is None:
            continue                                   # leyenda, notas, filas sueltas
        p = {}
        for c in campos:
            bruto = limpiar(fila[c["_col"]])
            p[c["key"]] = a_numero(bruto) if c["tipo"] == "numero" else bruto

        # El puesto puede ser 'ref' o 'X': se deja tal cual, la app ya lo contempla
        crudo = limpiar(fila[next(c["_col"] for c in campos if c["key"] == "puesto")])
        p["puesto"] = crudo if not isinstance(crudo, (int, float)) else a_numero(crudo)

        # id estable: si manana cambia el ranking, los favoritos siguen apuntando bien
        p["id"] = slug(p.get("firma")) + "--" + slug(p.get("plan"))

        m = re.match(r"\s*(\d+)\s*/\s*100", p.get("fiabilidad_por_que") or "")
        p["fiabilidad_nota"] = int(m.group(1)) if m else None

        planes.append(p)

    planes.sort(key=lambda x: (not isinstance(x["puesto"], int), x["puesto"]
                               if isinstance(x["puesto"], int) else 0))
    return campos, planes, subtitulo, leyenda


def leer_metodo(hoja):
    bloques, titulo = [], None
    for fila in hoja.iter_rows(values_only=True):
        a = limpiar(fila[0])
        b = limpiar(fila[1]) if len(fila) > 1 else None
        if a and b:
            bloques.append({"titulo": a, "texto": b})
        elif a and titulo is None:
            titulo = a
    return titulo, bloques


def leer_coste_real(hoja):
    m = a_numero(limpiar(hoja["C3"].value))
    return int(m) if m else 2


# --------------------------------------------------------------------------
# 4. Main
# --------------------------------------------------------------------------
def main():
    aqui = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="Excel de fondeadoras -> JSON para la app")
    ap.add_argument("--excel", default=str(aqui / "comparativa_fondeadoras.xlsx"))
    ap.add_argument("--salida", default=str(aqui / "fondeadoras.json"))
    ap.add_argument("--compacto", action="store_true", help="sin sangrado (mas pequeno)")
    args = ap.parse_args()

    ruta = Path(args.excel)
    if not ruta.exists():
        sys.exit("No encuentro el Excel: " + str(ruta))

    wb = openpyxl.load_workbook(ruta, data_only=True)
    if "Comparativa" not in wb.sheetnames:
        sys.exit("El libro no tiene hoja 'Comparativa'.")

    avisos = []
    campos, planes, subtitulo, leyenda = leer_comparativa(wb["Comparativa"], avisos)
    metodo_titulo, metodo = leer_metodo(wb["Metodo y avisos"]) if "Metodo y avisos" in wb.sheetnames else (None, [])
    meses = leer_coste_real(wb["Coste real"]) if "Coste real" in wb.sheetnames else 2

    esquema = []
    for c in campos:
        d = {k: v for k, v in c.items() if not k.startswith("_")}
        d["filtrable"] = bool(d.get("filtrable"))
        d["ordenable"] = bool(d.get("ordenable"))
        esquema.append(d)
    esquema.append(dict(CAMPO_NOTA))

    # Solo se exportan las claves declaradas: el JSON no arrastra restos
    claves = [c["key"] for c in esquema] + ["id"]
    planes = [{k: p.get(k) for k in claves} for p in planes]

    fechas = sorted({p["fecha_consulta"] for p in planes if p.get("fecha_consulta")})
    grupos = [g for g in ORDEN_GRUPOS if any(c["grupo"] == g for c in esquema)]
    grupos += [g for g in dict.fromkeys(c["grupo"] for c in esquema) if g not in grupos]
    nota_gen = limpiar(wb["Comparativa"]["A2"].value)

    salida = {
        "meta": {
            "esquema_version": 2,
            "titulo": limpiar(wb["Comparativa"]["A1"].value) or "Comparador de fondeadoras",
            "subtitulo": subtitulo if subtitulo and len(subtitulo) < 60 else "Cuenta de 50.000 USD",
            "n_planes": len(planes),
            "fecha_consulta": fechas[-1] if fechas else None,
            "fecha_consulta_min": fechas[0] if fechas else None,
            "generado": dt.datetime.now().strftime("%d/%m/%Y %H:%M"),
            "fuente": ruta.name,
            "grupos": grupos,
            "meses_default": meses,
            "coste_real_nota": COSTE_NOTA,
            "coste_real_aviso": COSTE_AVISO,
            "metodo_titulo": metodo_titulo or "Como se ha construido esto",
            "metodo": metodo,
            "nota_generacion": nota_gen,
            "leyenda": leyenda,
            "avisos_exportacion": avisos,
        },
        "campos": esquema,
        "planes": planes,
    }

    destino = Path(args.salida)
    destino.parent.mkdir(parents=True, exist_ok=True)
    with open(destino, "w", encoding="utf-8") as f:
        if args.compacto:
            json.dump(salida, f, ensure_ascii=False, separators=(",", ":"))
        else:
            json.dump(salida, f, ensure_ascii=False, indent=2)

    # ---------------- resumen ----------------
    print("OK  %s  (%.1f KB)" % (destino, destino.stat().st_size / 1024))
    print("    %d planes · %d campos (%d filtrables, %d ordenables)"
          % (len(planes), len(esquema),
             sum(1 for c in esquema if c["filtrable"]),
             sum(1 for c in esquema if c["ordenable"])))
    print("    fechas de consulta: %s a %s"
          % (salida["meta"]["fecha_consulta_min"], salida["meta"]["fecha_consulta"]))

    sin_precio = [p["id"] for p in planes if p.get("precio_tarifa") is None]
    if sin_precio:
        print("    AVISO  sin precio de tarifa: " + ", ".join(sin_precio))
    sin_nota = [p["id"] for p in planes if p.get("fiabilidad_nota") is None]
    if sin_nota:
        print("    AVISO  sin nota de fiabilidad: " + ", ".join(sin_nota))
    for a in avisos:
        print("    AVISO  " + a)

    hoy = dt.date.today()
    caducan = {}
    for p in planes:
        m = re.search(r"\d{4}-\d{2}-\d{2}", str(p.get("promo_caduca") or ""))
        if not m:
            continue
        dias = (dt.date.fromisoformat(m.group(0)) - hoy).days
        if dias <= 7:
            k = (p["firma"], m.group(0), dias)
            caducan[k] = caducan.get(k, 0) + 1
    for (firma, fecha, dias), n in sorted(caducan.items(), key=lambda x: x[0][2]):
        cuando = ("caducada hace %d dias" % -dias) if dias < 0 else (
            "caduca HOY" if dias == 0 else "caduca en %d dias" % dias)
        print("    AVISO  promo de %s (%s) %s — %d plan(es)" % (firma, fecha, cuando, n))


if __name__ == "__main__":
    main()
